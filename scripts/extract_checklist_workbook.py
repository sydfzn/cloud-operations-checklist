from __future__ import annotations

import json
import sys
from pathlib import Path
from openpyxl import load_workbook


def normalize(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def extract(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    records = []
    sheets = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        non_empty = [[normalize(value) for value in row] for row in rows if any(value is not None and str(value).strip() for value in row)]
        if not non_empty:
            continue
        header_index = 0
        for index, row in enumerate(non_empty[:12]):
            text = " ".join(value or "" for value in row).lower()
            if any(token in text for token in ("task", "activity", "check", "control", "owner", "status", "deliverable", "kyc", "transition", "onboard", "readiness")):
                header_index = index
                break
        headers = []
        for index, value in enumerate(non_empty[header_index]):
            headers.append(value or f"column_{index + 1}")
        seen = {}
        unique_headers = []
        for header in headers:
            count = seen.get(header, 0) + 1
            seen[header] = count
            unique_headers.append(header if count == 1 else f"{header}_{count}")
        data_rows = []
        for row in non_empty[header_index + 1:]:
            padded = row + [None] * max(0, len(unique_headers) - len(row))
            record = {unique_headers[index]: padded[index] for index in range(len(unique_headers)) if padded[index] not in (None, "")}
            if record:
                data_rows.append(record)
                records.append({"workbook": path.name, "worksheet": worksheet.title, **record})
        sheets.append({"worksheet": worksheet.title, "header_row": header_index + 1, "headers": unique_headers, "row_count": len(data_rows)})
    output = path.with_suffix(".normalized.json")
    output.write_text(json.dumps({"workbook": path.name, "sheets": sheets, "records": records}, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"workbook": path.name, "sheet_count": len(sheets), "sheets": sheets, "record_count": len(records), "normalized_file": str(output)}, indent=2))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: extract_checklist_workbook.py <xlsx>")
    extract(Path(sys.argv[1]))
